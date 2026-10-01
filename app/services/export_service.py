"""
Phase 7 — export. One function per format, all taking the same
report_service result shape plus a bit of metadata, so report_page.py
doesn't need format-specific logic.
"""
from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path

COMPANY_NAME = "Vgotel"


def export_csv(path: str, title: str, report: dict) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([title])
        w.writerow([f"Generated {dt.datetime.now().strftime('%d %b %Y %H:%M')}"])
        w.writerow([])
        w.writerow(report["columns"])
        for row in report["rows"]:
            w.writerow(row)
        if report.get("totals"):
            w.writerow([])
            for k, v in report["totals"].items():
                w.writerow([f"Total {k}", v])


def export_excel(path: str, title: str, report: dict, generated_by: str = "") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Report"

    ws.append([COMPANY_NAME])
    ws["A1"].font = Font(size=14, bold=True, italic=True)
    ws.append([title])
    ws["A2"].font = Font(size=12, bold=True)
    ws.append([f"Generated {dt.datetime.now().strftime('%d %b %Y %H:%M')}"
               + (f" by {generated_by}" if generated_by else "")])
    ws["A3"].font = Font(size=9, italic=True, color="6B7178")
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(report["columns"])
    for cell in ws[header_row]:
        cell.font = Font(bold=True)

    for row in report["rows"]:
        ws.append(row)

    if report.get("totals"):
        ws.append([])
        for k, v in report["totals"].items():
            ws.append([f"Total {k}", v])
            ws[f"A{ws.max_row}"].font = Font(bold=True)

    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    wb.save(path)


def export_pdf(path: str, title: str, report: dict, generated_by: str = "") -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
    )

    from reportlab.lib.enums import TA_LEFT

    styles = getSampleStyleSheet()
    company_style = styles["Title"].clone("CompanyName")
    company_style.alignment = TA_LEFT
    company_style.fontSize = 20

    doc = SimpleDocTemplate(
        path, pagesize=landscape(A4),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    story = []

    story.append(Paragraph(f"<b><i>{COMPANY_NAME}</i></b>", company_style))
    story.append(Paragraph(title, styles["Heading2"]))
    meta = f"Generated {dt.datetime.now().strftime('%d %b %Y %H:%M')}"
    if generated_by:
        meta += f" by {generated_by}"
    story.append(Paragraph(meta, styles["Normal"]))
    story.append(Spacer(1, 12))

    if report.get("summary"):
        summary_text = " &nbsp;&nbsp;|&nbsp;&nbsp; ".join(f"{k}: {v}" for k, v in report["summary"])
        story.append(Paragraph(summary_text, styles["Normal"]))
        story.append(Spacer(1, 10))

    table_data = [report["columns"]] + [[str(c) for c in row] for row in report["rows"]]
    if report.get("totals"):
        total_row = [""] * len(report["columns"])
        total_row[0] = "Total"
        for k, v in report["totals"].items():
            if k in report["columns"]:
                total_row[report["columns"].index(k)] = str(v)
        table_data.append(total_row)

    table = Table(table_data, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14171A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E7E5E0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FCFCFB")]),
    ]
    if report.get("totals"):
        style.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
        style.append(("LINEABOVE", (0, -1), (-1, -1), 1, colors.HexColor("#14171A")))
    table.setStyle(TableStyle(style))
    story.append(table)

    def _footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#A8A49B"))
        canvas.drawString(1.5 * cm, 1 * cm, f"{COMPANY_NAME} \u2014 Asset & Maintenance Register")
        canvas.drawRightString(landscape(A4)[0] - 1.5 * cm, 1 * cm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)


def default_export_path(report_key: str, fmt: str) -> str:
    out_dir = Path(__file__).resolve().parent.parent.parent / "exports"
    out_dir.mkdir(exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    return str(out_dir / f"{report_key}_{stamp}.{fmt}")
